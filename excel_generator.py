import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel(data, sections_wanted):
    """
    Génère un fichier Excel à partir des données extraites.
    """
    output = io.BytesIO()
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        workbook = writer.book
        
        # Styles
        header_fill = PatternFill(start_color="1a3a5c", end_color="1a3a5c", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        subheader_fill = PatternFill(start_color="2d6a9f", end_color="2d6a9f", fill_type="solid")
        subheader_font = Font(color="FFFFFF", bold=True, size=10)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_align = Alignment(horizontal='center', vertical='center')
        right_align = Alignment(horizontal='right', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')
        
        # Feuille 1: Informations générales
        if 'entreprise' in data:
            info_df = pd.DataFrame({
                'Champ': ['Raison Sociale', 'Identifiant Fiscal', 'Exercice', 'Date de Clôture'],
                'Valeur': [
                    data['entreprise'].get('raison_sociale', ''),
                    data['entreprise'].get('identifiant_fiscal', ''),
                    data['entreprise'].get('exercice', ''),
                    data['entreprise'].get('date_cloture', '')
                ]
            })
            info_df.to_excel(writer, sheet_name='Informations', index=False)
            worksheet = writer.sheets['Informations']
            
            # Appliquer le style
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
            worksheet.column_dimensions['A'].width = 25
            worksheet.column_dimensions['B'].width = 35
        
        # Feuille 2: Bilan Actif
        if 'bilan_actif' in sections_wanted and 'bilan_actif' in data:
            bilan_actif = data.get('bilan_actif', [])
            if bilan_actif:
                df_actif = pd.DataFrame(bilan_actif)
                df_actif.to_excel(writer, sheet_name='Bilan Actif', index=False)
                worksheet = writer.sheets['Bilan Actif']
                
                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center_align
                
                for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                    for cell in row:
                        cell.border = thin_border
                        if isinstance(cell.value, (int, float)):
                            cell.alignment = right_align
                        else:
                            cell.alignment = left_align
                
                for column in worksheet.columns:
                    max_length = 0
                    column = [cell for cell in column]
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = (max_length + 2)
                    worksheet.column_dimensions[get_column_letter(column[0].column)].width = adjusted_width
        
        # Feuille 3: Bilan Passif
        if 'bilan_passif' in sections_wanted and 'bilan_passif' in data:
            bilan_passif = data.get('bilan_passif', [])
            if bilan_passif:
                df_passif = pd.DataFrame(bilan_passif)
                df_passif.to_excel(writer, sheet_name='Bilan Passif', index=False)
                worksheet = writer.sheets['Bilan Passif']
                
                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center_align
                
                for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                    for cell in row:
                        cell.border = thin_border
                        if isinstance(cell.value, (int, float)):
                            cell.alignment = right_align
                        else:
                            cell.alignment = left_align
        
        # Feuille 4: CPC
        if 'cpc' in sections_wanted and 'cpc' in data:
            cpc_data = data.get('cpc', [])
            if cpc_data:
                df_cpc = pd.DataFrame(cpc_data)
                df_cpc.to_excel(writer, sheet_name='CPC', index=False)
                worksheet = writer.sheets['CPC']
                
                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center_align
                
                for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                    for cell in row:
                        cell.border = thin_border
                        if isinstance(cell.value, (int, float)):
                            cell.alignment = right_align
                        else:
                            cell.alignment = left_align
        
        # Feuille 5: Trésorerie
        if 'tresorerie' in sections_wanted and 'tresorerie' in data:
            treso = data.get('tresorerie', {})
            if treso:
                df_treso = pd.DataFrame({
                    'Poste': ['Actif Trésorerie', 'Passif Trésorerie', 'Trésorerie Nette'],
                    'Montant': [
                        treso.get('actif_tresorerie', 0),
                        treso.get('passif_tresorerie', 0),
                        treso.get('tresorerie_nette', 0)
                    ]
                })
                df_treso.to_excel(writer, sheet_name='Trésorerie', index=False)
                worksheet = writer.sheets['Trésorerie']
                
                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center_align
                
                for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                    for cell in row:
                        cell.border = thin_border
                        if isinstance(cell.value, (int, float)):
                            cell.alignment = right_align
                        else:
                            cell.alignment = left_align
        
        # Feuille 6: Ratios
        if 'ratios' in sections_wanted and 'ratios' in data:
            ratios = data.get('ratios', [])
            if ratios:
                df_ratios = pd.DataFrame(ratios)
                df_ratios.to_excel(writer, sheet_name='Ratios', index=False)
                worksheet = writer.sheets['Ratios']
                
                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center_align
                
                for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                    for cell in row:
                        cell.border = thin_border
                        cell.alignment = left_align
                
                worksheet.column_dimensions['A'].width = 30
                worksheet.column_dimensions['B'].width = 35
                worksheet.column_dimensions['C'].width = 15
                worksheet.column_dimensions['D'].width = 40
    
    output.seek(0)
    return output.getvalue()
